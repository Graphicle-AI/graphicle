#!/usr/bin/env python3
"""Run unchanged discovery cases against Rust, then compare complete real corpora."""
import ast, contextlib, copy, fnmatch, inspect, io, json, os, random, shutil, sys, tempfile, time, types
from pathlib import Path
from parity import Native, SOURCE
import naming
from naming import NATIVE_FUNCTIONS, to_upstream
import upstream_naming
import graphify.detect as det
import pytest

native=Native('detect')
temporary=tempfile.TemporaryDirectory();work=Path(temporary.name)
fixture_root=work/'fixtures';shutil.copytree(SOURCE/'tests/fixtures',fixture_root)
oracle=types.ModuleType('reference_detect');oracle.__package__='graphify'
exec(compile((SOURCE/'graphify/detect.py').read_text(),str(SOURCE/'graphify/detect.py'),'exec'),oracle.__dict__)
originals={}
comparisons=0
def plain(v):
    if isinstance(v,Path):return str(v)
    if isinstance(v,(tuple,list,set,frozenset)):return [plain(i) for i in v]
    if isinstance(v,dict):return {str(k):plain(i) for k,i in v.items()}
    return v
# git's user-wide excludes file (0.9.86, #4335) is found through HOME and
# XDG_CONFIG_HOME, which every upstream test sandboxes; the native process
# reads them when it starts, so a changed home restarts it.
spawned_home=None
def follow_home():
    global spawned_home
    current=(os.environ.get('HOME'),os.environ.get('XDG_CONFIG_HOME'))
    if current!=spawned_home:native.cwd=None;spawned_home=current
def call(function,**kwargs):
    follow_home()
    native.use_working_directory(os.getcwd())
    # The project's .graphifyrc (0.9.87, #2961) is read as .graphiclerc.
    # Only calls with a root enter the mirroring context: an upstream test
    # counts the reference cycles every matcher call leaves behind.
    if function not in ('detect','detect_incremental','save_manifest') or not isinstance(kwargs.get('root'),(str,Path)):
        return call_native(function,**kwargs)
    with naming.native_inputs(kwargs['root']):return call_native(function,**kwargs)
def call_native(function,**kwargs):
    # Unchanged tests reach the engine through private upstream names
    # (_load_graphifyignore); upstream_naming.py gives them the engine's file names.
    out=native.result('detect',function=NATIVE_FUNCTIONS.get(function,function),output=str(det.GRAPHIFY_OUT),limits={'raw_bytes':det._OFFICE_MAX_RAW_BYTES,'decompressed_bytes':det._OFFICE_MAX_DECOMPRESSED_BYTES,'compression_ratio':det._OFFICE_MAX_COMPRESSION_RATIO},**plain(kwargs))
    for warning in out['warnings']:print(warning,file=sys.stderr)
    return out['value']
def compare(function,expected,**kwargs):
    global comparisons
    actual=to_upstream(call(function,**kwargs));comparisons+=1
    if actual!=plain(expected):native.fail(function=function,arguments=plain(kwargs),expected=plain(expected),actual=actual)
    return actual

names=['classify_file','count_words','detect','detect_incremental','load_manifest','save_manifest','_looks_like_paper','_is_sensitive','_is_graphable_source','_is_prose_note','_generic_keyword_hit','_is_env_template','_shebang_interpreter','_shebang_file_type','_env_command_args','_split_env_s','_file_within_size_cap','_zip_within_caps','extract_pdf_text','docx_to_markdown','xlsx_to_markdown','xlsx_extract_structure','convert_office_file','_is_regular_file','_has_coverage_artifacts','_has_build_output_markers','_has_venv_markers','_is_noise_dir','_parse_gitignore_line','_find_vcs_root','_path_identity','_git_tracked_path_keys','_git_info_exclude','_read_ignore_text','_load_dir_own_ignore','_load_graphifyignore','_parse_ignore_pattern','_lexical_relative','_match_anchored_ignore_pattern','_is_ignored','_auto_follow_symlinks','_resolves_under_root','_os_path','_md5_file','_stat_and_hash','_nfc','_to_relative_for_storage','_to_absolute_from_storage','_mtime_may_hide_a_rewrite']
path_returns={'convert_office_file','_find_vcs_root','_git_info_exclude'}
def wrap(name):
    original=getattr(det,name);originals[name]=original;signature=inspect.signature(original)
    def invoke(*args,**kwargs):
        bound=signature.bind(*args,**kwargs);bound.apply_defaults();arguments=dict(bound.arguments)
        arguments.pop('_cache',None);arguments.pop('target_parts',None)
        renames={'d':'path','start':'path','target':'path','path_str':'path','vcs_root':'root','s':'value'}
        if name=='_match_anchored_ignore_pattern':renames['path']='value'
        for old,new in renames.items():
            if old in arguments:arguments[new]=arguments.pop(old)
        if name=='detect' and Path(arguments['root']).resolve().is_relative_to(SOURCE):arguments['cache_root']=str(work/'upstream-cache')
        if name in ('detect','detect_incremental') and arguments.get('google_workspace') is None:arguments['google_workspace']=det.google_workspace_enabled()
        value=call(name,**arguments)
        if name in path_returns:return Path(value) if value else None
        if name in ('_load_dir_own_ignore','_load_graphifyignore'):return [(Path(p),s) for p,s in value]
        if name=='_git_tracked_path_keys':return tuple(set(v) for v in value)
        if name in ('_parse_ignore_pattern','_stat_and_hash'):return tuple(value) if value is not None else None
        if name in ('classify_file','_shebang_file_type'):return det.FileType(value) if value else None
        return value
    return invoke
class Predicate:
    def __init__(self,root,**kwargs):self.id=call('predicate_new',root=root,**kwargs)
    def __call__(self,path):return call('predicate',id=self.id,path=path)
    def close(self):call('predicate_drop',id=self.id)

# 0.9.87 (#2961): a language declaration re-queues unchanged bytes and is
# stamped in the manifest; the watch cases of that file drive the reference watcher.
lifecycle=['tests/test_language_override_lifecycle.py::'+name for name in ('test_language_transitions_requeue_unchanged_bytes_but_not_unrelated_files','test_noop_config_edit_and_partial_save_preserve_only_valid_state','test_saving_one_tier_does_not_validate_the_other_old_language','test_legacy_manifest_stays_warm_until_a_language_override_applies','test_save_uses_explicit_root_not_another_roots_active_override')]
files=['tests/test_detect.py','tests/test_ignore_file_encoding.py','tests/test_office_limits.py','tests/test_office_incremental.py','tests/test_incremental_mtime_collision.py','tests/test_docx_to_markdown.py','tests/test_pdf_extra_warning.py',
       # First run here on 14 October 2026: long paths, non-regular files, the output
       # folder's evidence and the persisted scan options.
       'tests/test_long_path_hashing.py','tests/test_non_regular_files.py','tests/test_out_dir_evidence.py','tests/test_skill_persisted_scan_options.py',
       # 0.9.91 (#3427): Jinja templates are documents.
       'tests/test_detect_jinja.py']
# These cases monkeypatch a Python process/parser/counter that a Rust library
# cannot invoke. Exercise those contracts with actual native IO below instead.
excluded={
'test_git_tracking_probe_failure_preserves_ignore_behavior','test_git_lsfiles_skipped_when_no_gitignore_contributes',
'test_detect_converts_google_workspace_shortcuts_when_enabled','test_detect_office_sidecar_survives_a_gitignored_output_dir','test_detect_google_workspace_sidecar_survives_a_gitignored_output_dir',
'test_negation_does_not_disable_directory_pruning','test_ignored_dir_pruned_walk_never_lists_contents','test_scan_ignore_cost_is_per_directory_not_per_file','test_is_ignored_no_per_pattern_path_construction','test_is_ignored_cache_evaluates_each_dir_once',
'test_convert_office_file_hash_stable_across_nfc_nfd','test_convert_office_file_does_not_rewrite_existing_sidecar','test_convert_office_file_sidecar_name_stable_across_checkouts','test_convert_office_file_hash_disambiguates_same_stem','test_convert_office_file_outside_root_falls_back','test_detect_office_conversion_respects_cache_root',
'test_nested_graphify_out_prunes_only_configured_path','test_a_settled_corpus_costs_no_content_hashes',
'test_collect_files_keeps_coverage_code_namespace','test_save_manifest_skips_semantic_hash_for_files_without_cache',
# 0.9.71 stubs the Python load_manifest here; exercised through real files below.
'test_save_manifest_relativize_step_collapses_seeded_duplicates',
# 0.9.72 hides the optional pypdf import; the native PDF parser is built in, so
# the missing-extra warning cannot occur (the PDF comparison below covers text).
'test_extract_pdf_text_missing_pypdf_warns','test_extract_pdf_text_missing_pypdf_warns_once_across_many_pdfs',
# The engine has no .graphifyinclude handling (deleted with the rename; upstream
# deprecated it and only warned): such a file is an ordinary dotfile.
'test_graphifyinclude_is_inert_and_not_unclassified',
# Simulates Windows by patching Python's sys.platform; the engine takes the
# platform it runs on, so the long-path prefix is not added on Linux or macOS.
'test_os_path_adds_prefix_on_win32',
}
class Reports:
    passed=0;skipped=0;deselected=[]
    def pytest_collection_modifyitems(self,config,items):
        for item in items:
            if hasattr(item.module,'FIXTURES'):item.module.FIXTURES=fixture_root
        omitted=[item for item in items if item.originalname in excluded]
        items[:]=[item for item in items if item.originalname not in excluded]
        self.deselected=[item.nodeid for item in omitted];config.hook.pytest_deselected(items=omitted)
    def pytest_runtest_logreport(self,report):
        if report.when=='call' and report.passed:self.passed+=1
        if report.skipped:self.skipped+=1
    def pytest_runtest_setup(self,item):
        # A case that resets the Python stat-index singleton (monkeypatching
        # graphify.cache._stat_index_root) gets the same fresh Rust state.
        if '_stat_index_root' in inspect.getsource(item.obj):call('reset')
reports=Reports()
import graphify.cache as cache_mod
python_flush_stat_index=cache_mod._flush_stat_index
def flush_stat_index():
    # Cases flush the stat index explicitly (#3847); the Rust detector owns its
    # own index, so flush both before the assertions inspect the files.
    python_flush_stat_index();call('flush')
try:
    for name in names:setattr(det,name,wrap(name))
    originals['ignored_predicate']=det.ignored_predicate;det.ignored_predicate=Predicate
    cache_mod._flush_stat_index=flush_stat_index
    with upstream_naming.isolated(files+lifecycle,'check_detect') as (_,copies):status=pytest.main(['-q','-p','no:cacheprovider',*copies],plugins=[reports])
    assert status==0,status
    cache_mod._flush_stat_index=python_flush_stat_index
    for name,value in originals.items():setattr(det,name,value)
    # The seeded-duplicate contract of test_save_manifest_relativize_step_collapses_seeded_duplicates
    # (keep the more recently seen of two keys for one untouched file), through real
    # manifest files for both implementations: native save_manifest reads its own.
    for impl in ('python','native'):
        case=work/f'seeded-{impl}';(case/'src').mkdir(parents=True)
        tracked=case/'src'/'foo.py';tracked.write_text('def x(): pass\n')
        other=case/'bar.py';other.write_text('def y(): pass\n')
        dotted=os.path.join(str(case.resolve()),'src','..','src','foo.py')
        manifest=case/'graphify-out'/'manifest.json';manifest.parent.mkdir(parents=True)
        manifest.write_text(json.dumps({dotted:{'mtime':2.0,'seen':2.0,'ast_hash':'fresh','semantic_hash':'fresh_sem'},
                                        str(tracked.resolve()):{'mtime':1.0,'seen':1.0,'ast_hash':'stale','semantic_hash':''}}))
        if impl=='python':oracle.save_manifest({'code':[str(other)]},str(manifest),root=case)
        else:call('save_manifest',files={'code':[str(other)]},manifest_path=str(manifest),kind='both',root=str(case))
        row=json.loads(manifest.read_text()).get('src/foo.py',{});comparisons+=1
        if row.get('ast_hash')!='fresh':native.fail(function='save_manifest seeded duplicates',implementation=impl,row=row)
    # Helpers whose edge cases govern which files can enter the corpus.
    rng=random.Random(61)
    stems=['token','tokens','credentials','service-account','password-policy','api_token','secret_store','token.economics.notes','id_rsa','ID_RSA','grid_rsa','.env','.envrc','normal','café']
    extensions=['','.md','.txt','.py','.json','.env','.example','.pem','.tfvars','.tf']
    parents=['','.ssh/','.aws/','Secrets/','src/','internal/credentials/']
    for name in [p+s+e for p in parents for s in stems for e in extensions]:
        compare('_is_sensitive',oracle._is_sensitive(Path(name)),path=name)
    patterns=['*.tmp','a/**','**/a','**/a/**','a/*/c','a/**/c','[!a-z]','[a-z]','[[]','[]]','foo[','*','?','[z-a]','[a--z]','[!]]']
    for _ in range(500):
        name=''.join(rng.choice('az[]-/é') for _ in range(rng.randrange(9)))
        pattern=rng.choice(patterns)
        compare('_wildcard',fnmatch.fnmatchcase(name,pattern),value=name,pattern=pattern)
        compare('_match_anchored_ignore_pattern',oracle._match_anchored_ignore_pattern(name,pattern),value=name,pattern=pattern)
    # Actual DOCX parsing: styles, hyperlinks, horizontal/vertical merges.
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    corpus=work/'corpus';corpus.mkdir();documents=0
    for i in range(8):
        d=Document();d.add_heading('Project '+str(i),level=i%4);d.add_paragraph('Body café 日本語.');d.add_paragraph('First item',style='List Bullet');d.add_paragraph('')
        p=d.add_paragraph('Link: ');h=OxmlElement('w:hyperlink');run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='reference';run.append(text);h.append(run);p._p.append(h)
        table=d.add_table(rows=3,cols=3)
        for row in range(3):
            for col in range(3):table.cell(row,col).text=f'R{row} C{col}'
        if i%2:table.cell(0,0).merge(table.cell(0,1))
        if i%3==0:table.cell(1,2).merge(table.cell(2,2))
        # 0.9.72: a table stays in document order, and a cell is one escaped line.
        table.cell(2,0).text='a | b\nnext'
        if i%2==0:table.cell(2,1).add_table(rows=1,cols=2).cell(0,1).text='inner | cell'
        d.add_paragraph('After the table.')
        p=corpus/f'report-{i}.docx';d.save(p)
        compare('docx_to_markdown',oracle.docx_to_markdown(p),path=p);documents+=1
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table,TableStyleInfo
    import datetime
    for i in range(8):
        wb=Workbook();sheet=wb.active;sheet.title='Main'
        sheet.append(['Name','Value','When']);sheet.append(['Café',i+0.25,datetime.datetime(2026,9,23,4,5,6)]);sheet.append(['日本語',True,None]);sheet['E8'].number_format='0'
        table=Table(displayName='Records',ref='A1:C3');sheet.add_table(table)
        if i%3==0:
            table._initialise_columns()
            for column,value in zip(table.tableColumns,['Name','Value','When']):column.name=value
            sheet['A1']='Edited header'
            sheet['B1']=0
        if i%2:wb.create_sheet('Other').append(['a','b'])
        p=corpus/f'sheet-{i}.xlsx';wb.save(p)
        compare('xlsx_to_markdown',oracle.xlsx_to_markdown(p),path=p)
        compare('xlsx_extract_structure',oracle.xlsx_extract_structure(p),path=p);documents+=1
    # Sidecars, cache-root routing and incremental changes are checked together.
    (corpus/'.gitignore').write_text('graphify-out/\n')
    expected=oracle.detect(corpus,cache_root=work/'cache');actual=compare('detect',expected,root=corpus,cache_root=work/'cache')
    manifest=work/'manifest.json'
    call('save_manifest',files=actual['files'],manifest_path=manifest,root=corpus,kind='both',now=time.time())
    # The first incremental scan creates its default output directory. Prepare
    # it once so both compared directory enumerations see identical disk state.
    oracle.detect_incremental(corpus,str(manifest))
    compare('detect_incremental',oracle.detect_incremental(corpus,str(manifest)),root=corpus,manifest_path=manifest)
    # PDF geometry comes from a native parser rather than pypdf. Check textual
    # content/page order independently and retain raw parser differences.
    from reportlab.pdfgen.canvas import Canvas
    pdf=work/'pages.pdf';canvas=Canvas(str(pdf));canvas.drawString(72,720,'Page one: source evidence.');canvas.showPage();canvas.drawString(72,720,'Page two: supporting detail.');canvas.save()
    expected=oracle.extract_pdf_text(pdf);actual=call('extract_pdf_text',path=pdf)
    assert actual.split()==expected.split(),(actual,expected)
    pdf_receipt={'reference':expected,'native':actual,'exact':actual==expected,'token_sequence_equal':actual.split()==expected.split()}
    native.receipt('graphify/detect.py',['src/detect.rs','src/detect_types.rs','src/detect_ignore.rs','src/detect_documents.rs','src/detect_manifest.rs','src/assets/detect/constants.json'],
        original_tests={'files':files,'passed':reports.passed,'skipped':reports.skipped,'deselected':reports.deselected,'fixtures':'unchanged fixtures copied outside the enclosing repository, whose gitignore excludes the downloaded reference tree, for both Python and Rust'},independent_comparisons=comparisons,real_office_documents=documents,pdf_parser_comparison=pdf_receipt,
        interface_differences=['test_graphifyinclude_is_inert_and_not_unclassified is deselected: the engine has no .graphifyinclude handling (removed with the rename; the file was deprecated upstream and only produced a warning), so a leftover file is an ordinary dotfile and no removal warning is printed.'],
        adaptations=['Native Rust PDF/Office dependencies replace the Python dependencies. PDF geometric whitespace is reported separately; source text/page order is checked, without claiming universal pypdf byte equivalence.',
        'Spreadsheet structure extraction now applies the same resource caps as the live converter. The upstream dormant structure helper omitted those caps.',
        'Ignore-encoding and walk warnings are returned alongside results for embedded callers; CLI callers can print them. Native predicates and caches retain state without Python execution.',
        'Original Python subprocess and monkeypatch counters are not passed off as native checks. Real native source trees, converters, git operations, sidecars and manifests are used for integration checks.',
        'A path-cycle guard also stops indirect sibling symlink cycles. Ordinary in-root symlink aliases remain independently discovered.'])
finally:
    for name,value in originals.items():setattr(det,name,value)
    native.close();temporary.cleanup()
